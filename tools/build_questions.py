"""Build data/questions.json and js/checklist-data.js from the RTTT_and_ROC checklist.

The workbook has two audit question sets, chosen by store type:
  - ROC  (ROC_Audit tab)  -> SUPERSPAR / SPAR / SAVEMOR stores
  - RTTT (RTTT tab)       -> TOPS stores
Each question is scored 0-3, carries a Weight (%) and a built-in severity
(category: Critical / Major / Partial), grouped Section -> Focus Area.

Usage:  python3 tools/build_questions.py      (needs: pip install openpyxl)
"""
import openpyxl, json, os
from collections import OrderedDict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "samples", "RTTT_and_ROC_Checklist.xlsx")
wb = openpyxl.load_workbook(SRC, data_only=False)

def norm(v):
    if v is None: return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v

def sev_fix(v):
    if not v: return "Major"
    s = str(v).strip().lower()
    if s.startswith("cry") or s.startswith("cri"): return "Critical"  # "Crytal" typo -> Critical
    if s.startswith("maj"): return "Major"
    if s.startswith("part"): return "Partial"
    return str(v).strip().title()

def header_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(1, c).value)
        if h: m[h.lower()] = c
    return m

def col(hm, *starts):
    for key in hm:
        for s in starts:
            if key.startswith(s.lower()):
                return hm[key]
    return None

def parse_audit(sheet, audit_id, name, store_types):
    ws = wb[sheet]
    hm = header_map(ws)
    cSec = col(hm, "section"); cFocus = col(hm, "focus"); cQ = col(hm, "audit question")
    cW = col(hm, "weight"); cCat = col(hm, "category"); cMax = col(hm, "max score")
    cNote = col(hm, "auditor note")
    sections = OrderedDict()
    qn = 0
    for r in range(2, ws.max_row + 1):
        q = norm(ws.cell(r, cQ).value)
        if not q:
            continue
        qn += 1
        sec = norm(ws.cell(r, cSec).value) or "General"
        focus = norm(ws.cell(r, cFocus).value) or ""
        weight = ws.cell(r, cW).value if isinstance(ws.cell(r, cW).value, (int, float)) else 0
        maxsc = ws.cell(r, cMax).value if isinstance(ws.cell(r, cMax).value, (int, float)) else 3
        sev = sev_fix(ws.cell(r, cCat).value)
        note = norm(ws.cell(r, cNote).value) or ""
        sec_key = sec
        if sec_key not in sections:
            sections[sec_key] = []
        sections[sec_key].append({
            "id": audit_id.lower() + "-q" + str(qn),
            "ref": str(qn),
            "group": focus,          # Focus Area = sub-group within a section
            "text": q,
            "maxScore": int(maxsc),
            "weight": int(weight),
            "severity": sev,
            "guidance": note,
        })
    cats = []
    for i, (sec, qs) in enumerate(sections.items(), start=1):
        cats.append({
            "id": audit_id.lower() + "-cat-" + str(i),
            "number": str(i),
            "name": sec,
            "subtitle": "",
            "questions": qs,
        })
    return {
        "id": audit_id,
        "name": name,
        "storeTypes": store_types,
        "scoreMax": 3,
        "categories": cats,
    }

ROC = parse_audit("ROC_Audit", "ROC",
                  "ROC — Retail Operations Checklist",
                  ["SUPERSPAR", "SPAR", "SAVEMOR"])
RTTT = parse_audit("RTTT", "RTTT",
                   "RTTT — TOPS Store Audit",
                   ["TOPS"])

RATING_BANDS = [
    {"min": 95, "max": 100, "label": "Platinum"},
    {"min": 90, "max": 94, "label": "Gold"},
    {"min": 80, "max": 89, "label": "Silver"},
    {"min": 70, "max": 79, "label": "Bronze"},
    {"min": 0,  "max": 69, "label": "Improvement Required"},
]

STORE_TYPES = [
    {"label": "TOPS", "audit": "RTTT"},
    {"label": "SUPERSPAR", "audit": "ROC"},
    {"label": "SPAR", "audit": "ROC"},
    {"label": "SAVEMOR", "audit": "ROC"},
]

data = {
    "meta": {
        "title": "SPAR Retail Audit — ROC & RTTT",
        "shortTitle": "Retail Audit",
        "version": "V3 (27.08.26)",
        "scoreMax": 3,
        "passThresholdPct": 80,
        "scoreNote": "Score each question 0-3 (0 = not met, 3 = fully met), or N/A to exclude it. "
                     "Each question carries a Weight (%) and a built-in severity (Critical / Major / Partial). "
                     "Result % = sum(score/3 x weight) / sum(weight). Store passes at 80%.",
        "severityLevels": ["Critical", "Major", "Partial"],
        "remediationSeverities": ["Critical", "Major"],
        "ratingBands": RATING_BANDS,
        "source": "RTTT_and_ROC_Checklist.xlsx",
    },
    "storeTypes": STORE_TYPES,
    "audits": {"ROC": ROC, "RTTT": RTTT},
}

os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
os.makedirs(os.path.join(ROOT, "js"), exist_ok=True)
with open(os.path.join(ROOT, "data", "questions.json"), "w") as f:
    json.dump(data, f, indent=2, ensure_ascii=False)
with open(os.path.join(ROOT, "js", "checklist-data.js"), "w") as f:
    f.write("// AUTO-GENERATED from data/questions.json by tools/build_questions.py — do not edit by hand.\n"
            "window.CHECKLIST = " + json.dumps(data, ensure_ascii=False) + ";\n")

# report
for aud in (ROC, RTTT):
    tw = sum(q["weight"] for c in aud["categories"] for q in c["questions"])
    nq = sum(len(c["questions"]) for c in aud["categories"])
    sev = {}
    for c in aud["categories"]:
        for q in c["questions"]:
            sev[q["severity"]] = sev.get(q["severity"], 0) + 1
    print(f"{aud['id']:5} storeTypes={aud['storeTypes']}")
    print(f"      sections={len(aud['categories'])} questions={nq} weight%sum={tw} severity={sev}")
